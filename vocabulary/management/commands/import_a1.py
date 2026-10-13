from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from vocabulary.models import (
    VocabularyGroup,
    VocabularyLesson,
    VocabularyLevel,
    VocabularyWord,
)


class Command(BaseCommand):
    help = "Import A1 vocabulary bank from Excel"

    def add_arguments(self, parser):
        parser.add_argument(
            "excel_file",
            type=str,
            help="Path to A1 Excel file",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        file_path = Path(options["excel_file"])

        if not file_path.exists():
            raise CommandError(
                f"File not found: {file_path}"
            )

        self.stdout.write(
            f"Reading file: {file_path}"
        )

        workbook = load_workbook(
            file_path,
            data_only=True,
        )

        sheet_name = "A1_Master"

        if sheet_name not in workbook.sheetnames:
            raise CommandError(
                f'Sheet "{sheet_name}" not found.'
            )

        worksheet = workbook[sheet_name]

        # -----------------------------------------
        # Read headers from row 3
        # -----------------------------------------

        headers = {}

        for column_number, cell in enumerate(
            worksheet[3],
            start=1,
        ):
            if cell.value:
                headers[str(cell.value).strip()] = column_number

        required_columns = [
            "Source ID",
            "Original headword",
            "Normalized word",
            "Sense",
            "POS (source)",
            "Group #",
            "Group",
            "Lesson in group",
            "Global lesson",
            "Item in lesson",
            "Translation RU",
            "Translation KZ",
            "Example EN",
            "Example RU",
            "Example KZ",
        ]

        missing_columns = [
            column
            for column in required_columns
            if column not in headers
        ]

        if missing_columns:
            raise CommandError(
                "Missing columns: "
                + ", ".join(missing_columns)
            )

        # -----------------------------------------
        # Create A1 level
        # -----------------------------------------

        level, created = VocabularyLevel.objects.update_or_create(
            code="A1",
            defaults={
                "name": "Beginner",
                "order": 1,
                "is_active": True,
            },
        )

        if created:
            self.stdout.write(
                self.style.SUCCESS(
                    "Created level A1"
                )
            )
        else:
            self.stdout.write(
                "Level A1 already exists"
            )

        created_groups = 0
        created_lessons = 0
        created_words = 0
        updated_words = 0
        skipped_rows = 0

        # Cache so we don't hit database unnecessarily
        groups_cache = {}
        lessons_cache = {}

        # -----------------------------------------
        # Import rows
        # Data begins at row 4
        # -----------------------------------------

        for row_number in range(
            4,
            worksheet.max_row + 1,
        ):

            def value(column_name):
                column_number = headers[column_name]

                result = worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).value

                if isinstance(result, str):
                    return result.strip()

                return result

            source_id = value("Source ID")

            # Ignore empty rows
            if source_id is None:
                skipped_rows += 1
                continue

            word = value("Normalized word")

            if not word:
                self.stdout.write(
                    self.style.WARNING(
                        f"Row {row_number}: no word, skipped"
                    )
                )
                skipped_rows += 1
                continue

            group_number = value("Group #")
            group_name = value("Group")

            lesson_number = value(
                "Lesson in group"
            )

            global_lesson = value(
                "Global lesson"
            )

            item_number = value(
                "Item in lesson"
            )

            # -------------------------------------
            # Group
            # -------------------------------------

            group_number = int(group_number)

            if group_number not in groups_cache:

                group_code = (
                    f"G{group_number:02d}"
                )

                group, group_created = (
                    VocabularyGroup.objects.update_or_create(
                        level=level,
                        code=group_code,
                        defaults={
                            "name": group_name,
                            "order": group_number,
                        },
                    )
                )

                groups_cache[group_number] = group

                if group_created:
                    created_groups += 1

            group = groups_cache[group_number]

            # -------------------------------------
            # Lesson
            # -------------------------------------

            lesson_number = int(
                lesson_number
            )

            lesson_key = (
                group.id,
                lesson_number,
            )

            if lesson_key not in lessons_cache:

                lesson, lesson_created = (
                    VocabularyLesson.objects.update_or_create(
                        group=group,
                        number=lesson_number,
                        defaults={
                            "global_number": int(global_lesson),
                            "title": f"Lesson {lesson_number}",
                            "order": lesson_number,
                            "is_active": True,
                        },
                    )
                )

                lessons_cache[lesson_key] = lesson

                if lesson_created:
                    created_lessons += 1

            lesson = lessons_cache[
                lesson_key
            ]

            # -------------------------------------
            # Word
            # -------------------------------------

            word_object, word_created = (
                VocabularyWord.objects.update_or_create(
                    source_id=int(source_id),
                    defaults={
                        "lesson": lesson,
                        "original_headword":
                            value("Original headword") or "",
                        "word": word,
                        "sense":
                            value("Sense") or "",
                        "part_of_speech":
                            value("POS (source)") or "",
                        "translation_ru":
                            value("Translation RU") or "",
                        "translation_kz":
                            value("Translation KZ") or "",
                        "example_en":
                            value("Example EN") or "",
                        "example_ru":
                            value("Example RU") or "",
                        "example_kz":
                            value("Example KZ") or "",
                        "order":
                            int(item_number),
                        "is_active": True,
                    },
                )
            )

            if word_created:
                created_words += 1
            else:
                updated_words += 1

        # -----------------------------------------
        # Final statistics
        # -----------------------------------------

        total_groups = (
            VocabularyGroup.objects
            .filter(level=level)
            .count()
        )

        total_lessons = (
            VocabularyLesson.objects
            .filter(group__level=level)
            .count()
        )

        total_words = (
            VocabularyWord.objects
            .filter(lesson__group__level=level)
            .count()
        )

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                "A1 import completed!"
            )
        )

        self.stdout.write(
            f"Groups created: {created_groups}"
        )

        self.stdout.write(
            f"Lessons created: {created_lessons}"
        )

        self.stdout.write(
            f"Words created: {created_words}"
        )

        self.stdout.write(
            f"Words updated: {updated_words}"
        )

        self.stdout.write(
            f"Rows skipped: {skipped_rows}"
        )

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                "Database totals:"
            )
        )

        self.stdout.write(
            f"A1 groups: {total_groups}"
        )

        self.stdout.write(
            f"A1 lessons: {total_lessons}"
        )

        self.stdout.write(
            f"A1 words: {total_words}"
        )